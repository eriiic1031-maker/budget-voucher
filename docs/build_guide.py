"""테스트 가이드 Word 파일 생성 (docs/예산및전표관리_테스트가이드.docx).

화면 이미지는 docs/images/*.png (빨간 박스·번호 표시된 캡처)를 사용한다.
    .venv\\Scripts\\python docs\\build_guide.py
"""
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = Path(__file__).parent
IMG = HERE / "images"
OUT = HERE / "예산및전표관리_테스트가이드.docx"
FONT = "맑은 고딕"
RED = RGBColor(0xE5, 0x39, 0x35)
BLUE = RGBColor(0x2F, 0x5B, 0xD8)
GRAY = RGBColor(0x6B, 0x75, 0x88)
CIRCLED = "①②③④⑤⑥⑦⑧⑨"

doc = Document()

# ---------- 기본 서식 ----------
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
sec.orientation = WD_ORIENT.PORTRAIT
for side in ("left_margin", "right_margin"):
    setattr(sec, side, Cm(2.0))
sec.top_margin, sec.bottom_margin = Cm(2.0), Cm(1.8)
CONTENT_W = Cm(17.0)


def set_font(style_or_run, size=None, bold=None, color=None):
    f = style_or_run.font
    f.name = FONT
    rpr = style_or_run.element.get_or_add_rPr() if hasattr(style_or_run, "element") else None
    if rpr is not None:
        fonts = rpr.find(qn("w:rFonts"))
        if fonts is None:
            fonts = OxmlElement("w:rFonts")
            rpr.append(fonts)
        for k in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
            fonts.set(qn(k), FONT)
    if size:
        f.size = Pt(size)
    if bold is not None:
        f.bold = bold
    if color is not None:
        f.color.rgb = color


for name, size in [("Normal", 10), ("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11.5), ("Title", 24)]:
    st = doc.styles[name]
    set_font(st, size=size)
    if name.startswith("Heading") or name == "Title":
        st.font.bold = True
        st.font.color.rgb = RGBColor(0x1B, 0x23, 0x33)
doc.styles["Normal"].paragraph_format.space_after = Pt(4)
doc.styles["Normal"].paragraph_format.line_spacing = 1.12

# 바닥글: 쪽 번호
fp = sec.footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = fp.add_run()
for tag, text in (("begin", None), (None, "PAGE"), ("end", None)):
    if tag:
        el = OxmlElement("w:fldChar")
        el.set(qn("w:fldCharType"), tag)
    else:
        el = OxmlElement("w:instrText")
        el.set(qn("xml:space"), "preserve")
        el.text = text
    r._r.append(el)
set_font(r, size=9, color=GRAY)


# ---------- 도우미 ----------
def para(text="", size=None, bold=False, color=None, align=None, after=None):
    p = doc.add_paragraph()
    if text:
        add_rich(p, text, size=size, bold=bold, color=color)
    if align:
        p.alignment = align
    if after is not None:
        p.paragraph_format.space_after = Pt(after)
    return p


EMOJI = re.compile("([🌀-🫿☀-➿⭐✅☑][️]?)")


def add_rich(p, text, size=None, bold=False, color=None):
    """**굵게** 표기를 지원하는 간단한 런 추가. 이모지는 Segoe UI Emoji 로 (맑은 고딕에는 없어서 네모로 보임)."""
    parts = text.split("**")
    for i, part in enumerate(parts):
        for seg in EMOJI.split(part):
            if not seg:
                continue
            run = p.add_run(seg)
            set_font(run, size=size, bold=bold or (i % 2 == 1), color=color)
            if EMOJI.fullmatch(seg):
                run.font.name = "Segoe UI Emoji"
                run.element.rPr.rFonts.set(qn("w:eastAsia"), "Segoe UI Emoji")
                run.element.rPr.rFonts.set(qn("w:ascii"), "Segoe UI Emoji")
                run.element.rPr.rFonts.set(qn("w:hAnsi"), "Segoe UI Emoji")
    return p


def shade(cell, hex_fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcpr.append(shd)


def table(headers, rows, widths, header_fill="EEF2F8", first_col_red=False, font_size=9.5):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = ""
        add_rich(c.paragraphs[0], h, size=font_size, bold=True)
        c.paragraphs[0].paragraph_format.space_after = Pt(0)
        shade(c, header_fill)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            p = cells[i].paragraphs[0]
            if i == 0 and first_col_red:
                add_rich(p, v, size=12, bold=True, color=RED)
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                add_rich(p, v, size=font_size)
            p.paragraph_format.space_after = Pt(0)
    for row in t.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def steps(items):
    """번호 목록 (Word 기본 번호 스타일)."""
    for it in items:
        p = doc.add_paragraph(style="List Number")
        add_rich(p, it)
        p.paragraph_format.space_after = Pt(2)


def bullets(items):
    for it in items:
        p = doc.add_paragraph(style="List Bullet")
        add_rich(p, it)
        p.paragraph_format.space_after = Pt(2)


def screen(file, caption, marks):
    """캡처 이미지 + 번호 설명표. marks: [(위치, 할 일·기대 결과)]"""
    from PIL import Image
    w, h = Image.open(IMG / file).size
    max_h = Cm(13.0)
    if CONTENT_W * h / w > max_h:          # 너무 긴 캡처는 높이 기준으로 줄인다
        doc.add_picture(str(IMG / file), height=max_h)
    else:
        doc.add_picture(str(IMG / file), width=CONTENT_W)
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.paragraphs[-1].paragraph_format.keep_with_next = True
    para(caption, size=8.5, color=GRAY, align=WD_ALIGN_PARAGRAPH.CENTER, after=6)
    table(["번호", "위치", "할 일 · 기대 결과"],
          [(CIRCLED[i], pos, desc) for i, (pos, desc) in enumerate(marks)],
          widths=[1.3, 3.7, 12.0], first_col_red=True)


def scenario(title, first=False):
    """시나리오는 새 쪽에서 시작 (첫 번째는 장 제목 바로 아래)."""
    h = doc.add_heading(title, level=2)
    h.paragraph_format.page_break_before = not first
    h.paragraph_format.keep_with_next = True
    return h


def page_break():
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


# ================================================================ 표지 ====
para("", after=60)
t = doc.add_paragraph(style="Title")
add_rich(t, "예산 및 전표관리", size=28, bold=True)
para("테스트 가이드", size=18, color=BLUE, after=24)
para("총무 담당자용 샘플 프로그램을 화면 순서대로 눌러 보며 확인하는 문서입니다. "
     "화면 그림의 **빨간 박스 번호**와 아래 설명표의 **①②③** 번호가 서로 맞습니다.", size=11, after=30)
table(["항목", "내용"], [
    ("기준일", "2026-10-21 (10월은 21일까지 데이터)"),
    ("예시 데이터", "2025-01 ~ 2026-10 세금계산서 1,695건 (월 70~80건) + 연간 항목 10종"),
    ("작년(2025) 비용", "공급가액 합계 약 121.7억원"),
    ("예산", "2026년 = 2025년 같은 달 실적 × 1.05 · 2025년 = 2026년 ÷ 1.05 (총무 예산 계정만)"),
    ("SAP 연결", "모의 모드 — 실제 SAP에는 아무것도 전송되지 않음"),
    ("로그인", "없음 — 바로 총무 담당자 화면으로 열림"),
    ("프로그램 폴더", "D:\\총무 Ai\\Budget_Voucher"),
], widths=[4.0, 13.0])
para("업체명·사업자번호·사람 이름은 모두 가상의 값입니다.", size=9, color=GRAY)
page_break()

# ============================================================ 1. 시작하기 ====
doc.add_heading("1. 시작하기", level=1)
para("PowerShell(터미널)을 **D:\\총무 Ai\\Budget_Voucher** 폴더에서 열고 아래 순서대로 실행합니다.")
table(["순서", "명령", "설명"], [
    ("1", ".venv\\Scripts\\python -m voucher.seed",
     "예시 데이터 만들기·초기화. 약 10초 뒤 '세금계산서 1,695건 · 121.7억원'이 나오면 성공. "
     "테스트 중 바꾼 내용을 처음으로 되돌릴 때도 이 명령"),
    ("2", ".venv\\Scripts\\python -m streamlit run streamlit_app.py",
     "프로그램 켜기 → 브라우저에서 http://localhost:8501"),
    ("3", "Ctrl + C", "프로그램 끄기"),
], widths=[1.3, 7.2, 8.5], font_size=9)
para("처음 설치하는 PC라면 1번 전에 `python -m venv .venv` 와 "
     "`.venv\\Scripts\\python -m pip install -r requirements.txt` 를 한 번 실행합니다.", size=9, color=GRAY)

doc.add_heading("화면 보기 전환 (로그인 대신)", level=2)
para("왼쪽 사이드바 **화면 보기**에서 사람을 고르면 그 사람의 권한으로 화면이 바뀝니다. "
     "**기준 월**은 그 아래에서 고릅니다 (기본 2026-10).")
table(["화면 보기", "권한", "부서", "보이는 메뉴"], [
    ("총무 담당자", "관리자", "총무팀", "전체 9개"),
    ("홍길동", "사용자", "영업1팀", "5개 (총무 예산·홈택스·기준정보·사용자 관리 없음)"),
    ("김철수", "사용자", "물류팀", "5개"),
], widths=[3.2, 2.2, 2.4, 9.2])

doc.add_heading("메뉴 한눈에 보기", level=2)
table(["메뉴", "하는 일", "누가"], [
    ("대시보드", "오늘 할 일, 총무 예산 경고, 이번 달 지표·차트, 최근 활동", "모두 (사용자는 우리 부서 기준)"),
    ("총무 예산", "월 예산 대비 사용·월말 예상, 초과(예상) 경고와 이유, 연간 누적, 예산 입력", "관리자"),
    ("전표 점검", "정기 전표(기한·금액) / 비정기 전표(내용 확인)", "모두"),
    ("검토·전기", "자동 분류가 확신 못 한 건 수정 후 SAP 전기, 승인 요청 처리, 일괄 전기", "관리자 (사용자는 '검토·승인 요청')"),
    ("부서별 비용", "부서에 배부된 총무 비용 (계정별·사용자별)", "모두"),
    ("세금계산서 조회", "월별 전체 목록, SAP 전표 내용, 처리 이력", "모두"),
    ("홈택스 대사", "홈택스 엑셀로 메일 안 온 세금계산서 찾기·자동 등록", "관리자"),
    ("기준정보", "정기 전표, 계정 규칙, 배부 기준, 차량, 공급업체, 부서, 계정(예산 주관)", "관리자"),
    ("사용자 관리", "사람·권한·부서 등록 ('화면 보기' 목록에 반영)", "관리자"),
], widths=[3.0, 9.5, 4.5], font_size=9)

# ======================================================== 2. 테스트 시나리오 ====
doc.add_heading("2. 테스트 시나리오", level=1).paragraph_format.page_break_before = True
para("아래 숫자는 예시 데이터를 **방금 만든 상태**(기준 월 2026-10) 기준입니다. T1부터 순서대로 진행하면 앞 단계 결과가 "
     "뒤 단계로 이어집니다. 숫자가 다르면 `python -m voucher.seed` 로 초기화한 뒤 다시 시작하세요.")

scenario("T1. 대시보드 (총무 담당자)", first=True)
screen("01_dashboard.png", "그림 1. 대시보드 — 처음 열면 나오는 화면", [
    ("사이드바 · 화면 보기", "'총무 담당자 · 관리자 (총무팀)'이 선택되어 있는지 확인"),
    ("사이드바 · 기준 월", "2026-10"),
    ("오늘 할 일", "**승인 요청 1 · 검토 대기 1 · 전기 대기 5 · 예산 경고 6 · 정기 전표 확인 3 · 비정기 미확인 18 · "
                 "연간 항목 1**. **카드 아무 곳이나 누르면** 해당 화면으로 이동 (예: '검토 대기' → 검토·전기 화면에서 "
                 "'검토 대기'가 선택된 채 열림, '연간 항목' → 전표 점검의 연간 항목 탭)"),
    ("총무 예산", "수선비·차량유지비·렌탈료·수도광열비 **초과** 카드와 이유 두 줄. 아래 차트의 계정 목록에 "
                "운반비·광고선전비·교육훈련비·지급수수료가 **없어야** 함 (총무 예산이 아님)"),
])

scenario("T1-2. 작년보다 늘어난 비용 · 달력")
para("대시보드를 아래로 내리면 나옵니다. 총무 예산 계정만, 작년 **같은 날짜까지**와 비교합니다.")
screen("13_yoy.png", "그림 1-2. 작년보다 늘어난 비용", [
    ("비교 기간", "**이번 달** / **올해 누적** 전환"),
    ("증가 표", "이번 달: 수선비 +4,013만원 (원인: 튼튼설비 '대회의실 리모델링 3,800만원'), 차량유지비 +1,335만원 (엔진 수리·주유), "
              "복리후생비 (창립기념 행사 1,500만원 · **사내 피트니스 제휴 신규 300만원**) 등. "
              "'올해 누적'으로 바꾸면 **보안관제 서비스 신규**, 본사 정기주차 증가(7월 5면 추가), 전기요금 인상이 보임"),
])
screen("14_calendar.png", "그림 1-3. 이번 달 달력 (28일 선택)", [
    ("접기·펼치기", "제목 줄을 누르면 달력이 접히고 다시 누르면 펼쳐짐"),
    ("색 구분", "파랑 정기 전표 기한 · 보라 연간 항목 · 주황 작년 이 시기 지출 · 초록 세무·마감 일정 · 빨강 공휴일"),
    ("날짜 선택", "**28일**을 누르면 테두리가 생기고 아래에 그날 일정이 나옴. 오늘(21일)은 파란 동그라미"),
    ("그날 일정", "경비·보안관제·시설관리·청소 용역 기한(수신 대기) + 작년 10/28 전기 정기 안전점검 150만원"),
    ("다음 달 미리 보기", "다음 달 연간 항목: 승강기 정기검사 (11/30, 예상 68만원)"),
])
para("확인할 것: 26일 '부가세 신고 마감'(25일이 일요일이라 다음 영업일), 30일 '월말 결산 마감', 3·5·9일 공휴일, "
     "31일 '전기 정기 안전점검'(연간 항목).", size=9.5)

scenario("T2. 총무 예산 — 경고와 이유")
screen("02_budget_month.png", "그림 2. 총무 예산 › 이번 달", [
    ("보기 전환", "이번 달 / 연간 누적 / 예산 입력"),
    ("이번 달 지표", "예산 · 사용 · 월말 예상 · **초과 4 / 초과 예상 2**"),
    ("예산 경고 · 왜 그런가요?", "수선비: '대회의실 리모델링 3,800만원 (예산의 422%)'. 다른 카드도 확인 — "
                            "차량유지비: 엔진 수리 748만원·지출 속도 40%↑ / 렌탈료: 복합기 1,019만→1,256만원 / "
                            "복리후생비(초과 예상): 구내식당 1억 4,995만원 예정 + 창립기념 행사 1,500만원"),
])
para("아래로 내리면 **주의 3건**(건물관리비·용역비·지급임차료) 펼침과 항목별 현황표, 예산·월말 예상·작년 비교 차트가 있습니다. "
     "'연간 누적'에서는 누적 예산·올해·작년 선 3개와 계정별 집행률을 봅니다. "
     "사이드바 기준 월을 **2025년**으로 바꾸면 2025년 예산(2026년 예산 ÷ 1.05, 연 111.3억원)으로 같은 화면을 볼 수 있습니다.",
     size=9.5)
screen("03_budget_input.png", "그림 3. 총무 예산 › 예산 입력", [
    ("연도", "2026"),
    ("월별 예산 표", "수선비 행의 **10월**을 `60000000` 으로 바꾸기"),
    ("저장", "**저장** → '이번 달'로 돌아가면 수선비가 경고에서 빠지고 **초과 3건**. "
           "테스트 뒤 `9000000` 으로 되돌리기 (또는 seed로 초기화)"),
])

scenario("T3. 전표 점검 — 정기 전표")
screen("04_recurring.png", "그림 4. 전표 점검 › 정기 전표 (복합기 렌탈료 행 선택)", [
    ("구분", "정기 전표 · 확인 필요 8 / 비정기 전표 · 미확인 18"),
    ("정기 전표 지표", "**완료 10/19 · 기한 경과 1 · 금액 상이 1 · 수신·미전기 1 · 수신 대기 6** "
                    "(올해 새로 계약한 사내 피트니스 제휴·보안관제 서비스 포함)"),
    ("정기 전표 목록", "🔴 본사 임차료(기한 경과), 🟠 복합기 렌탈료(금액 상이), 🟡 정수기·공기청정기(수신·미전기). "
                   "**복합기 렌탈료** 행 왼쪽 체크칸을 눌러 선택"),
    ("확인 처리", "'확인 완료' + 메모 \"초과 출력 요금, 공급자 확인\" → **저장** → ☑️ 수동 확인, 완료 11/19. "
              "메모 없이 저장하면 '메모를 입력하세요'가 나와야 함"),
])

scenario("T4. 전표 점검 — 비정기 전표")
screen("05_nonrecurring.png", "그림 5. 전표 점검 › 비정기 전표 (3건 선택)", [
    ("비정기 지표", "**비정기 50건 · 미확인 18 · 큰 금액(300만원↑) 9 · 신규 거래처 2**"),
    ("필터", "미확인만 보기 / 총무 예산만. '총무 예산만'을 끄면 택배·광고·교육 건이 추가로 보임"),
    ("비정기 목록", "💰 = 300만원 이상, 🆕 = 최근 3개월 신규 거래처. 행 3개의 체크칸을 눌러 선택"),
    ("확인 완료", "메모(선택) 입력 → **3건 확인 완료** → 미확인 **15**, 선택한 행은 목록에서 사라짐"),
])

scenario("T4-2. 전표 점검 — 연간 항목")
para("세금계산서가 오지 않는 재산세·자동차세·보험 갱신·법정 점검처럼 **매년 이 시기에 처리할 항목**을 체크합니다. "
     "처리 전에는 예상 금액이 총무 예산의 월말 예상에 들어가고, 처리하면 실적으로 바뀝니다.")
screen("15_annual.png", "그림 5-2. 전표 점검 › 연간 항목", [
    ("구분", "**연간 항목 · 미처리 1** (대시보드 '연간 항목' 카드를 눌러도 이 화면)"),
    ("이번 달 항목", "**전기 정기 안전점검** (기한 10/31, 예상 165만원). 처리 금액·처리일·메모 입력 → **처리 완료** → "
                 "✅ 처리 완료, 미처리 0. '처리 취소'로 되돌릴 수 있음"),
    ("전체 일정", "1~12월 표: 재산세(건물분) 7월 · 재산세(토지분) 9월 · 자동차세 6·12월 · 주민세 8월 · 보험 갱신 1·3월 · "
              "정기 소독 분기 등. 항목 추가는 기준정보 › 연간 항목"),
])
para("확인할 것: 처리 완료 후 총무 예산 › 이번 달의 용역비 이유에서 '아직 처리 안 한 연간 항목' 문장이 사라짐. "
     "기준 월을 2026-07로 바꾸면 재산세(건물분) 1,961만원이 '처리 완료'로 보임.", size=9.5)

scenario("T5. 사용자 화면 — 승인 요청 (홍길동)")
screen("06_user_dashboard.png", "그림 6. 홍길동(영업1팀)으로 본 대시보드", [
    ("화면 보기", "**홍길동 · 사용자 (영업1팀)** 으로 변경"),
    ("메뉴", "5개만 보임 (총무 예산·홈택스 대사·기준정보·사용자 관리 없음)"),
    ("우리 부서 지표", "제목이 '영업1팀 · 총무 예산 항목 이번 달' — 영업1팀에 배부된 총무 비용만 집계"),
])
screen("07_user_review.png", "그림 7. 검토·승인 요청 › 검토 필요 (새봄문구 선택)", [
    ("확인이 필요한 이유", "'SAP 공급업체 미등록 … 계정 규칙 없음' — 처음 보는 거래처라 자동 분류 불가"),
    ("차변 분개 표", "**계정** 칸을 더블클릭해 '소모품비', **부서** 칸에서 '영업1팀' 선택. "
                 "SAP 공급업체 코드는 몰라도 비워 둠"),
    ("요청 메모 · 승인 요청", "메모 입력 → **관리자에게 승인 요청** → '내 승인 요청' **2건**. "
                         "계정·부서가 비어 있으면 버튼이 눌리지 않아야 함"),
])

scenario("T6. 검토·전기 (총무 담당자)")
screen("08_admin_review.png", "그림 8. 검토·전기 › 승인 요청 (동네주유소 선택)", [
    ("구분", "승인 요청 / 검토 대기 / 전기 대기. 화면 보기를 **총무 담당자**로 되돌린 뒤 '승인 요청' 선택"),
    ("요청 내용", "홍길동 님의 요청 메모 확인"),
    ("SAP 공급업체 코드", "`200066` 입력 (이 사업자번호를 마스터에 저장 ✔)"),
    ("SAP 전기 · 반려", "**SAP 전기** → 'SAP 전표 51… 생성' 알림, 목록에서 사라짐. 코드 없이 누르면 막혀야 함. "
                    "새봄문구는 **반려** → '검토 대기'로 돌아가는지 확인"),
])
para("이어서 '전기 대기 5' → **5건 일괄 전기** → 0건. 전표 점검에서 정수기·공기청정기가 ✅ 전기 완료로 바뀝니다.", size=9.5)

scenario("T7. 홈택스 대사 (무료 엑셀)")
screen("09_hometax.png", "그림 9. 홈택스 대사 › 엑셀 올리기 후", [
    ("탭", "**홈택스 엑셀 올리기 (무료)**"),
    ("파일", "`samples\\홈택스_매입목록_예시_202610.xlsx` 를 끌어다 놓기 → **가져오기**"),
    ("결과", "'63건 확인 · 시스템에 없던 2건 중 2건 등록' (10월 본사 임차료, 바른인쇄 명함)"),
    ("대사 결과", "홈택스 매입 63 · **일치 63** · 미수집 0 · 금액 불일치 0 · 홈택스 미확인 0"),
])
para("확인할 것: 전표 점검에서 본사 임차료가 🔴 기한 경과 → 🟡 수신·미전기로 바뀌고, 검토·전기의 전기 대기에 들어와 있음. "
     "같은 파일을 한 번 더 올리면 '등록 0건' (중복 등록 없음).", size=9.5)

scenario("T8. 부서별 비용 · 세금계산서 조회")
screen("10_dept_costs.png", "그림 10. 부서별 비용 (영업1팀)", [
    ("부서 선택", "**영업1팀** (관리자는 모든 부서 선택 가능, 사용자는 자기 부서만)"),
    ("부서 지표", "이번 달 비용 · 건수 · 가장 큰 항목, 전월 대비"),
    ("계정별·사용자별", "계정별 막대와 사용자(차량 담당자)별 막대, 아래 상세 내역"),
])
screen("11_invoices.png", "그림 11. 세금계산서 조회 (전기 완료 건 선택)", [
    ("세금계산서 목록", "아무 '전기 완료' 행 선택"),
    ("SAP 전표", "차변(비용·부서별)·대변(공급업체)·부가세대급금 라인과 처리 이력이 보임"),
])

scenario("T9. 기준정보 (선택)")
screen("12_masters.png", "그림 12. 기준정보 › 계정", [
    ("탭", "**계정** 탭 선택"),
    ("예산 주관", "운반비(82400000)의 예산 주관을 `총무` 로 변경"),
    ("저장", "**저장** → 대시보드 차트와 총무 예산에 운반비가 나타남 (예산 미편성이라 ⚪ 예산 없음). "
           "테스트 뒤 `물류팀` 으로 되돌리기"),
])

# ========================================================== 3. 체크리스트 ====
doc.add_heading("3. 테스트 체크리스트", level=1).paragraph_format.page_break_before = True
para("결과 칸에 ○ / △ / × 를 적고, 문제가 있으면 메모에 본 화면을 적어 주세요.")
table(["#", "확인 항목", "결과", "메모"], [
    ("T1", "대시보드 할 일 7장 · 카드 눌러 이동 · 예산 카드 · 타 부서 계정 제외", "", ""),
    ("T1-2", "작년보다 늘어난 비용(이번 달/올해 누적) · 달력 접기·날짜 선택", "", ""),
    ("T2", "예산 경고 4/2 · 이유 문장 · 예산 수정 반영", "", ""),
    ("T3", "정기 전표 10/19 · 복합기 수동 확인", "", ""),
    ("T4", "비정기 50건 · 3건 확인 → 미확인 15", "", ""),
    ("T4-2", "연간 항목 처리 완료 → 미처리 0 · 예산 이유 반영", "", ""),
    ("T5", "홍길동 화면 · 새봄문구 승인 요청", "", ""),
    ("T6", "승인 → SAP 전기 · 반려 · 일괄 전기", "", ""),
    ("T7", "홈택스 엑셀 → 2건 등록 · 전부 일치 · 재업로드 0건", "", ""),
    ("T8", "부서별 비용 · SAP 전표 내용 보기", "", ""),
    ("T9", "예산 주관 변경 → 운반비 표시", "", ""),
], widths=[1.3, 8.7, 1.8, 5.2])

# ======================================================= 4. 문제 해결·한계 ====
doc.add_heading("4. 문제 해결과 알려진 한계", level=1)
table(["증상", "해결"], [
    ("숫자가 가이드와 다름", "앞 테스트의 전기·확인이 남아 있는 것. `.venv\\Scripts\\python -m voucher.seed` 로 초기화"),
    ("'사용자가 없습니다' 화면", "예시 데이터가 없는 상태. 위 seed 명령 실행"),
    ("포트 8501을 이미 사용 중", "켜져 있는 다른 창을 끄거나 실행 명령 끝에 `--server.port 8502`"),
    ("화면 보기를 바꾼 뒤 메뉴가 사라짐", "사용자 권한에는 관리자 메뉴가 없는 것이 정상"),
    ("홈택스 엑셀을 못 읽음", "오류에 '인식된 컬럼'이 나오니 그 내용을 알려 주세요"),
], widths=[5.0, 12.0])
para("알려진 한계 (샘플이라 일부러 비워 둔 것)", bold=True)
bullets([
    "로그인이 없고 누구나 '화면 보기'로 관리자가 될 수 있음. 운영 전에는 사내 로그인(SSO)을 붙여야 함",
    "SAP는 모의 모드라 전표번호만 만들고 실제 전송은 하지 않음",
    "메일 자동 수집은 메일 계정 설정이 필요해 이 가이드에서는 제외",
    "월말 예상·경고 기준(초과 예상 102%, 주의 98%, 큰 건 예산 5%·300만원)은 임시 값 — 테스트 후 조정",
    "예산은 회사 전체 계정 단위. 부서별 예산은 아직 없음",
])

doc.save(OUT)
print("saved", OUT)
